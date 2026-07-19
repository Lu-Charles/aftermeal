"""Read-only source checks. Requires openpyxl and Pillow, separate from the demo.

Run with the existing spreadsheet-capable Python environment. This never writes
to the research directory and never infers corrected mass labels from images.
"""
import argparse
import hashlib
import json
from pathlib import Path
import openpyxl
from PIL import Image, ImageDraw, ImageOps

def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--research-root', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True, help='New directory for evidence')
    parser.add_argument('--ids', nargs='+', type=int, default=[133, 400, 117, 476, 81, 82, 113, 142])
    args = parser.parse_args()
    research = args.research_root.resolve()
    if args.output.exists() or args.output.resolve().is_relative_to(research):
        raise ValueError('Choose a new output directory outside the research workspace.')
    root = Path(__file__).resolve().parent.parent
    workbook_path = research / 'lefood/LeFood-Set Leftovers Food Dataset/LeFood-Set/data_original.xlsx'
    original = json.loads((research / 'results/lefood_records.json').read_text())
    packaged = {r['record_id']: r for r in json.loads((root / 'data/records.json').read_text())}
    hashes = json.loads((research / 'results/image_hashes.json').read_text())
    images = list((research / 'lefood').rglob('*.JPG'))
    paths = {p.name: p for p in images}
    if len(paths) != len(images):
        raise ValueError('Ambiguous duplicate image filenames.')
    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        if len(workbook.worksheets) != 1:
            raise ValueError('Expected the original single-sheet LeFood workbook.')
        sheet = workbook.worksheets[0]
        headers = next(sheet.iter_rows(values_only=True))
        expected = ('ID', 'Name of the food', 'Image Before Eaten', 'Weight Before Eaten (g)', 'Image After Eaten', 'Weight After Eaten (g)', 'Visual Estimation by Observer (1-7)')
        if headers != expected:
            raise ValueError('Source columns differ from the expected LeFood schema.')
        source_rows = {}
        for number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            key = values[0]
            if key in source_rows:
                raise ValueError(f'Duplicate source ID: {key}')
            source_rows[key] = (number, values)
        evidence = []
        checked_hashes = {}
        for row in original:
            number, values = source_rows[row['source_id']]
            record_id = f"L{row['source_id']}"
            actual = (row['source_id'], row['food'], row['before_image'], row['before_g'], row['after_image'], row['after_g'], row['observer_level'])
            if actual != values or number != row['excel_row']:
                raise ValueError(f'Imported values or workbook row differ for {record_id}')
            record = packaged[record_id]
            if abs(values[5] / values[3] - record['fraction']) > 1e-09:
                raise ValueError(f'Packaged ratio differs from source weights for {record_id}')
            if record['group'] != 'L:' + values[1]:
                raise ValueError(f'Packaged group differs for {record_id}')
            photo_evidence = {}
            for kind in ('before', 'after'):
                name = row[f'{kind}_image']
                if name not in checked_hashes:
                    checked_hashes[name] = digest(paths[name])
                if checked_hashes[name] != hashes[name]:
                    raise ValueError(f'Raw image changed since original extraction: {name}')
                photo_evidence[kind] = {'filename': name, 'sha256': checked_hashes[name]}
            if row['source_id'] in args.ids:
                evidence.append({'record_id': record_id, 'food': row['food'], 'sheet': sheet.title, 'range': f'A{number}:G{number}', 'before_g': values[3], 'after_g': values[5], 'recorded_fraction': record['fraction'], 'images': photo_evidence})
        expected_ids = {r['record_id'] for r in packaged.values() if r['source'] == 'LeFood'}
        if {f"L{r['source_id']}" for r in original} != expected_ids:
            raise ValueError('Research and packaged LeFood record sets differ.')
        if {e['record_id'] for e in evidence} != {f'L{i}' for i in args.ids}:
            raise ValueError('A requested review ID is absent from the packaged data.')
        evidence.sort(key=lambda e: args.ids.index(int(e['record_id'][1:])))
        report = {'scope': 'Source mapping and byte identity only; not validation of physical mass or visual labels.', 'workbook': str(workbook_path), 'workbook_sha256': digest(workbook_path), 'source_rows': len(source_rows), 'matched_packaged_pairs': len(original), 'checked_raw_images': len(checked_hashes), 'review_selection': 'Explicit diagnostic IDs; not a random sample.', 'review_pairs': evidence}
    finally:
        workbook.close()
    args.output.mkdir(parents=True)
    (args.output / 'source-audit.json').write_text(json.dumps(report, indent=2, allow_nan=False) + '\n')
    canvas = Image.new('RGB', (900, 260 * len(evidence)), 'white')
    draw = ImageDraw.Draw(canvas)
    for j, row in enumerate(evidence):
        label = f"{row['record_id']} {row['food']} | before {row['before_g']}g | after {row['after_g']}g ({100 * row['recorded_fraction']:.2f}%)"
        draw.text((8, 260 * j + 8), label, fill='black')
        for k, kind in enumerate(('before', 'after')):
            with Image.open(paths[row['images'][kind]['filename']]) as source:
                photo = ImageOps.contain(ImageOps.exif_transpose(source).convert('RGB'), (440, 220))
            canvas.paste(photo, (450 * k, 260 * j + 32))
    canvas.save(args.output / 'review-pairs.jpg')
    print(json.dumps({k: report[k] for k in ('source_rows', 'matched_packaged_pairs', 'checked_raw_images')}, indent=2))
if __name__ == '__main__':
    main()
